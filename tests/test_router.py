import pytest
from openpyxl import load_workbook

from conftest import ROOT


def excel_test_questions():
    wb = load_workbook(ROOT / "data" / "AI_Agent_Workflow_Assessment.xlsx", read_only=True)
    rows = list(wb["Test_Questions"].iter_rows(values_only=True))[1:]
    return [(r[0], r[1]) for r in rows if r[0]]


@pytest.mark.parametrize("wid,request_text", excel_test_questions())
def test_excel_test_questions_route_correctly(agent, wid, request_text):
    d = agent.router.route(request_text, agent.settings.min_route_confidence)
    assert d.workflow_id == wid, d.reasoning


@pytest.mark.parametrize("wid,request_text", [
    ("WF001", "what's low in stock?"),
    ("WF002", "compare our prices with the supplier"),
    ("WF003", "clean up the supplier upload file"),
    ("WF006", "are there any repeated SKUs?"),
    ("WF008", "what search intent do these terms have"),
    ("WF010", "give me a report on workflow errors"),
])
def test_paraphrases(agent, wid, request_text):
    assert agent.router.route(request_text, agent.settings.min_route_confidence).workflow_id == wid


@pytest.mark.parametrize("request_text", ["Book me a flight to Paris", "hello", "Show me the product list"])
def test_out_of_scope_requests_ask_to_clarify(agent, request_text):
    state = agent.run(request_text)
    assert state["workflow_id"] is None and state["status"] == "no_workflow"
    assert "I can help with" in state["answer"]
